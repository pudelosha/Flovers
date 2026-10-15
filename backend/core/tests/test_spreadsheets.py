from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook
from core.spreadsheets import append_safe_row


def test_untrusted_strings_round_trip_as_literal_text_without_formulas_or_error_cells():
    workbook = Workbook()
    worksheet = workbook.active
    values = ["=1+1", "=HYPERLINK(\"https://example.invalid\",\"click\")", "+SUM(1,2)", "@SUM(1,2)", "#DIV/0!", "Ordinary text"]
    append_safe_row(worksheet, values)
    stream = BytesIO()
    workbook.save(stream)
    loaded = load_workbook(BytesIO(stream.getvalue()), data_only=False).active
    assert [cell.value for cell in loaded[1]] == values
    assert all(cell.data_type == "s" for cell in loaded[1])


def test_xml_invalid_controls_are_removed_and_scalar_types_are_preserved():
    workbook = Workbook()
    worksheet = workbook.active
    append_safe_row(worksheet, ["A\x00\x01\ud800\ufffeB\n\tC", 12, 2.5, True, None, date(2026, 10, 15)])
    stream = BytesIO()
    workbook.save(stream)
    loaded = load_workbook(BytesIO(stream.getvalue())).active
    assert loaded["A1"].value == "AB\n\tC"
    assert [loaded.cell(1, column).value for column in range(2, 6)] == [12, 2.5, True, None]
    assert loaded["F1"].value.date() == date(2026, 10, 15)
