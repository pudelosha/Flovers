from datetime import timedelta
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from locations.models import Location
from plant_instances.models import PlantInstance
from reminders.models import Reminder, ReminderTask
from reminders.serializers import ReminderSerializer
from reminders.views import ReminderDetailView

pytestmark = pytest.mark.django_db
User = get_user_model()


@pytest.fixture
def care():
    user = User.objects.create_user(email="care@example.com", password="password123")
    room_a = Location.objects.create(user=user, name="Old room")
    room_b = Location.objects.create(user=user, name="New room")
    plant_a = PlantInstance.objects.create(user=user, location=room_a, display_name="Old plant")
    plant_b = PlantInstance.objects.create(user=user, location=room_b, display_name="New plant")
    reminder = Reminder.objects.create(user=user, plant=plant_a, type="water",
                                       start_date=timezone.localdate(), interval_value=7)
    client = APIClient()
    client.force_authenticate(user)
    return SimpleNamespace(user=user, reminder=reminder, plant_a=plant_a, plant_b=plant_b, client=client)


def test_prevalidated_stale_patches_preserve_other_committed_fields(care):
    request = SimpleNamespace(user=care.user)
    first = ReminderSerializer(Reminder.objects.get(pk=care.reminder.pk), data={"interval_value": 14},
                               partial=True, context={"request": request})
    second = ReminderSerializer(Reminder.objects.get(pk=care.reminder.pk), data={"is_active": False},
                                partial=True, context={"request": request})
    assert first.is_valid(), first.errors
    assert second.is_valid(), second.errors
    view = ReminderDetailView()
    view.request = request
    view.perform_update(first)
    view.perform_update(second)
    care.reminder.refresh_from_db()
    assert care.reminder.interval_value == 14
    assert not care.reminder.is_active
    assert not care.reminder.tasks.filter(status="pending").exists()


def test_prefetched_task_uses_latest_saved_schedule(care):
    task = care.reminder.ensure_one_pending_task()
    stale = ReminderTask.objects.select_related("reminder").get(pk=task.pk)
    response = care.client.patch(reverse("reminder-detail", args=[care.reminder.pk]),
                                 {"interval_value": 14}, format="json")
    assert response.status_code == 200
    next_task = stale.mark_complete_and_spawn_next(note="actual care")
    assert next_task.due_date == timezone.localdate() + timedelta(days=14)
    assert stale.completed_event["interval_value"] == 14


def test_replayed_completion_cannot_change_original_date_note_or_spawn_twice(care):
    task = care.reminder.ensure_one_pending_task()
    replay = ReminderTask.objects.select_related("reminder").get(pk=task.pk)
    task.mark_complete_and_spawn_next(note="original", completed_date=timezone.localdate() - timedelta(days=1))
    original_at = task.completed_at
    assert replay.mark_complete_and_spawn_next(note="replacement") is None
    replay.refresh_from_db()
    assert replay.note == "original"
    assert replay.completed_at == original_at
    assert care.reminder.tasks.filter(status="pending").count() == 1


def test_paused_or_removed_task_is_not_completed_from_stale_instance(care):
    task = care.reminder.ensure_one_pending_task()
    response = care.client.patch(reverse("reminder-detail", args=[care.reminder.pk]), {"is_active": False}, format="json")
    assert response.status_code == 200
    with pytest.raises(ReminderTask.DoesNotExist):
        task.mark_complete_and_spawn_next(note="after pause")
    assert not care.reminder.tasks.exists()


def test_pending_constraint_and_cancelled_visibility(care):
    care.reminder.ensure_one_pending_task()
    with pytest.raises(IntegrityError), transaction.atomic():
        ReminderTask.objects.create(reminder=care.reminder, user=care.user, due_date=timezone.localdate())
    cancelled = ReminderTask.objects.create(reminder=care.reminder, user=care.user,
                                            due_date=timezone.localdate(), status="cancelled")
    response = care.client.get(reverse("reminder-tasks-list"))
    assert cancelled.pk not in [row["id"] for row in response.json()]


def complete_then_move(care):
    task = care.reminder.ensure_one_pending_task()
    task.mark_complete_and_spawn_next(note="Watered old plant")
    event = dict(task.completed_event)
    response = care.client.patch(reverse("reminder-detail", args=[care.reminder.pk]), {
        "plant": care.plant_b.pk, "type": "fertilize", "interval_value": 14,
    }, format="json")
    assert response.status_code == 200, response.json()
    task.refresh_from_db()
    assert task.completed_event == event
    return task


def test_completed_journal_and_task_snapshot_keep_original_plant_and_type(care):
    task = complete_then_move(care)
    old = care.client.get(reverse("plant-instance-journal", args=[care.plant_a.pk]))
    new = care.client.get(reverse("plant-instance-journal", args=[care.plant_b.pk]))
    assert old.status_code == 200 and new.status_code == 200
    history = old.json()
    assert len(history) == 1
    assert history[0].pop("completed_by")["id"] == str(care.user.pk)
    assert history[0].pop("access")["is_owner"] is True
    assert history == [{"id": task.pk, "type": "water", "completed_at": history[0]["completed_at"], "note": "Watered old plant"}]
    assert new.json() == []
    row = care.client.get(reverse("reminder-tasks-list"), {"status": "completed"}).json()[0]
    assert row["completed_event"]["plant_id"] == care.plant_a.pk
    assert row["completed_event"]["type"] == "water"


def test_completed_snapshot_is_preserved_on_later_save(care):
    task = complete_then_move(care)
    original = dict(task.completed_event)
    task.completed_event = {"plant_id": care.plant_b.pk, "type": "fertilize"}
    task.note = "Updated note"
    task.save()
    task.refresh_from_db()
    assert task.completed_event == original
    assert task.note == "Updated note"


@pytest.mark.parametrize("filters", [{"plantId": "original"}, {"location": "Old room"}, {"types": ["watering"]}])
def test_completed_export_filters_and_cells_use_original_event(care, filters):
    complete_then_move(care)
    if filters.get("plantId") == "original":
        filters = {"plantId": care.plant_a.pk}
    with patch("reminders.views.send_templated_email") as send:
        response = care.client.post(reverse("reminder-task-export-email"), filters, format="json")
    assert response.status_code == 200
    kwargs = send.call_args.kwargs
    assert kwargs["context"]["completed_count"] == 1
    wb = load_workbook(BytesIO(kwargs["attachments"][0]["content"]))
    assert list(wb.active.values)[1][2:5] == ("Watering", "Old plant", "Old room")
    assert list(wb.active.values)[1][9] == "7 days"


def test_export_note_is_literal_text_and_invalid_xml_controls_are_removed(care):
    task = care.reminder.ensure_one_pending_task()
    task.mark_complete_and_spawn_next(note="=1+1\x00")
    with patch("reminders.views.send_templated_email") as send:
        response = care.client.post(reverse("reminder-task-export-email"), {}, format="json")
    assert response.status_code == 200
    wb = load_workbook(BytesIO(send.call_args.kwargs["attachments"][0]["content"]))
    assert wb.active["H2"].value == "=1+1"
    assert wb.active["H2"].data_type == "s"


@pytest.mark.parametrize("filters", [{"mode": "plant", "plantId": "original"}, {"mode": "location", "location": "Old room"}, {"mode": "types", "types": ["watering"]}])
def test_completed_bulk_delete_filters_original_event(care, filters):
    task = complete_then_move(care)
    if filters.get("plantId") == "original":
        filters = {"mode": "plant", "plantId": str(care.plant_a.pk)}
    response = care.client.post(reverse("reminder-task-bulk-delete"), filters, format="json")
    assert response.status_code == 200 and response.json()["deleted"] == 1
    assert not ReminderTask.objects.filter(pk=task.pk).exists()
    assert care.reminder.tasks.filter(status="pending").count() == 1
